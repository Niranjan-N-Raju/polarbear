from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).resolve().parents[1] / "Polar-Bear-Application-Overview.pptx"

NAVY = "14283F"
INK = "1D2B3A"
TEAL = "19A6B8"
ORANGE = "F07832"
YELLOW = "E8B849"
PAPER = "F5F7F4"
WHITE = "FFFFFF"
MIST = "EAF5F5"
PEACH = "FFF0E6"
LINE = "DCE3E1"
SLATE = "586777"
MUTED = "7A8792"
FONT = "Aptos"
DISPLAY = "Aptos Display"


def color(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(0.8)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, shade=LINE, width=1.2):
    shape = slide.shapes.add_connector(
        1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(shade)
    shape.line.width = Pt(width)
    return shape


def text(slide, value, x, y, w, h, size=14, shade=INK, bold=False,
         font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(shade)
    return box


def bullet_list(slide, items, x, y, w, line_height=0.46, size=14, shade=SLATE,
                accent=TEAL):
    for index, item in enumerate(items):
        yy = y + index * line_height
        rect(slide, x, yy + 0.10, 0.08, 0.08, accent, radius=True)
        text(slide, item, x + 0.22, yy, w - 0.22, line_height - 0.03,
             size=size, shade=shade)


def base_slide(pres, background=PAPER):
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(background)
    return slide


def footer(slide, number, dark=False):
    shade = "A9BAC8" if dark else MUTED
    line(slide, 0.58, 7.03, 12.75, 7.03, "354A60" if dark else LINE, 0.7)
    text(slide, "POLAR BEAR  /  STORE OPERATIONS", 0.60, 7.12, 5.2, 0.18,
         size=8, shade=shade, bold=True)
    text(slide, f"APPLICATION OVERVIEW  |  {number:02d}", 9.40, 7.12, 3.35, 0.18,
         size=8, shade=shade, bold=True, align=PP_ALIGN.RIGHT)


def heading(slide, eyebrow, title, subtitle=None):
    text(slide, eyebrow.upper(), 0.68, 0.50, 11.8, 0.24, size=9,
         shade=ORANGE, bold=True)
    text(slide, title, 0.68, 0.88, 12.0, 0.58, size=27,
         shade=NAVY, bold=True, font=DISPLAY)
    if subtitle:
        text(slide, subtitle, 0.70, 1.58, 11.8, 0.42, size=13, shade=SLATE)


def label(slide, value, x, y, w, fill=MIST, shade=TEAL):
    rect(slide, x, y, w, 0.30, fill, radius=True)
    text(slide, value.upper(), x + 0.12, y + 0.055, w - 0.24, 0.16,
         size=8, shade=shade, bold=True, valign=MSO_ANCHOR.MIDDLE)


def panel(slide, x, y, w, h, title, body, accent=TEAL, index=None):
    rect(slide, x, y, w, h, WHITE, line=LINE, radius=True)
    rect(slide, x, y, 0.055, h, accent)
    tx = x + 0.28
    if index:
        text(slide, index, tx, y + 0.24, 0.52, 0.4, size=21,
             shade=accent, bold=True, font=DISPLAY)
        tx += 0.58
    text(slide, title, tx, y + 0.24, w - (tx - x) - 0.22, 0.30,
         size=15, shade=NAVY, bold=True)
    text(slide, body, x + 0.28, y + 0.78, w - 0.54, h - 0.96,
         size=11, shade=SLATE)


def add_line_items(slide, items, x, y, w, accent=TEAL, gap=0.52):
    for i, (title, detail) in enumerate(items):
        yy = y + i * gap
        rect(slide, x, yy + 0.03, 0.10, 0.10, accent, radius=True)
        text(slide, title, x + 0.22, yy, w - 0.22, 0.24, size=12,
             shade=NAVY, bold=True)
        text(slide, detail, x + 0.22, yy + 0.25, w - 0.22, 0.28,
             size=9, shade=SLATE)


def make_deck():
    pres = Presentation()
    pres.slide_width = Inches(13.333)
    pres.slide_height = Inches(7.5)
    pres.core_properties.title = "Polar Bear Store Operations | Application Overview"
    pres.core_properties.subject = "Executive overview of the Polar Bear store operations application"
    pres.core_properties.author = "Polar Bear"
    pres.core_properties.keywords = "Polar Bear, store operations, application overview"

    # 1. Cover
    slide = base_slide(pres, NAVY)
    rect(slide, 0.72, 0.72, 0.72, 0.72, TEAL, radius=True)
    text(slide, "PB", 0.72, 0.91, 0.72, 0.24, size=20, shade=WHITE,
         bold=True, font=DISPLAY, align=PP_ALIGN.CENTER)
    text(slide, "POLAR BEAR  /  APPLICATION OVERVIEW", 0.72, 1.86, 7.7, 0.25,
         size=10, shade="B7C6D3", bold=True)
    text(slide, "Store operations,\nconnected.", 0.72, 2.45, 7.7, 1.48,
         size=35, shade=WHITE, bold=True, font=DISPLAY)
    text(slide, "A practical workspace for central oversight and\nconsistent day-to-day execution across stores.",
         0.76, 4.18, 6.6, 0.77, size=16, shade="D0DAE2")
    label(slide, "For Polar Bear officials", 0.76, 5.35, 2.48,
          fill="243C54", shade="C8E9ED")
    # Simple operating stack
    rect(slide, 8.42, 1.28, 4.08, 4.94, "1D344B", line="39536A", radius=True)
    text(slide, "ONE WORKSPACE", 8.78, 1.66, 3.2, 0.25, size=9,
         shade="B7C6D3", bold=True)
    stack = [
        ("01", "SALES", "Daily store transactions", TEAL),
        ("02", "PEOPLE", "Employees, attendance, pay", ORANGE),
        ("03", "STOCK", "Catalog, expiry, inventory", YELLOW),
        ("04", "FINANCE", "Expenses and P&L", TEAL),
    ]
    for i, (n, title, detail, accent) in enumerate(stack):
        yy = 2.16 + i * 0.92
        text(slide, n, 8.80, yy, 0.38, 0.28, size=10, shade=accent, bold=True)
        text(slide, title, 9.30, yy - 0.01, 2.62, 0.25, size=13,
             shade=WHITE, bold=True)
        text(slide, detail, 9.30, yy + 0.32, 2.70, 0.25, size=9, shade="C2CFD9")
        if i < 3:
            line(slide, 8.83, yy + 0.64, 12.12, yy + 0.64, "39536A", 0.7)
    footer(slide, 1, dark=True)

    # 2. Executive framing
    slide = base_slide(pres)
    heading(slide, "Executive summary", "One operating model. Two operating levels.",
            "Central teams oversee the network; store teams manage the work in their assigned store.")
    panel(slide, 0.70, 2.40, 5.82, 3.52, "Network oversight",
          "Admin and sub-admin views bring store and user administration together, with a dashboard designed for network-level visibility.",
          accent=TEAL, index="A")
    add_line_items(slide, [
        ("Stores", "Create and manage store records"),
        ("Users", "Assign roles and store access"),
        ("Dashboard", "Review store and workforce counts"),
    ], 1.02, 3.48, 4.98, accent=TEAL, gap=0.65)
    panel(slide, 6.78, 2.40, 5.82, 3.52, "Store execution",
          "Store admins work within their assigned store across sales, inventory, workforce, and financial workflows.",
          accent=ORANGE, index="B")
    add_line_items(slide, [
        ("Daily operations", "Sales, inventory, withdrawals"),
        ("People", "Employees, attendance, salary"),
        ("Performance", "Store-level trends and P&L"),
    ], 7.10, 3.48, 4.98, accent=ORANGE, gap=0.65)
    footer(slide, 2)

    # 3. Roles
    slide = base_slide(pres)
    heading(slide, "Role-based access", "Clear responsibility at every level.",
            "The experience adapts to the signed-in role and authorized store scope.")
    rect(slide, 0.70, 2.36, 5.78, 3.88, WHITE, line=LINE, radius=True)
    label(slide, "Admin / sub-admin", 1.02, 2.70, 1.84, fill=MIST, shade=TEAL)
    text(slide, "Set up the network", 1.02, 3.22, 4.9, 0.42, size=20,
         shade=NAVY, bold=True, font=DISPLAY)
    bullet_list(slide, [
        "Create and maintain store records",
        "Manage authorized user accounts",
        "Review network-level dashboard summaries",
        "Navigate to stores and user administration",
    ], 1.05, 3.95, 4.98, line_height=0.48, size=12, accent=TEAL)
    rect(slide, 6.84, 2.36, 5.78, 3.88, WHITE, line=LINE, radius=True)
    label(slide, "Store admin", 7.16, 2.70, 1.32, fill=PEACH, shade=ORANGE)
    text(slide, "Run the store", 7.16, 3.22, 4.9, 0.42, size=20,
         shade=NAVY, bold=True, font=DISPLAY)
    bullet_list(slide, [
        "Record sales and operating expenses",
        "Manage stock and expiry workflows",
        "Maintain employee and attendance records",
        "Review payroll and store financials",
    ], 7.19, 3.95, 4.98, line_height=0.48, size=12, accent=ORANGE)
    footer(slide, 3)

    # 4. Operating loop
    slide = base_slide(pres)
    heading(slide, "Daily operating rhythm", "The application follows the workday.",
            "Capture activity once, then use it across the store's operational views.")
    steps = [
        ("01", "Capture", "Record sales, payment mix, and expenses.", TEAL),
        ("02", "Control", "Receive stock and track expiry by product.", ORANGE),
        ("03", "Coordinate", "Maintain employee and attendance records.", YELLOW),
        ("04", "Review", "Use P&L, salary, and dashboard summaries.", TEAL),
    ]
    for i, (number, title, detail, accent) in enumerate(steps):
        x = 0.72 + i * 3.16
        rect(slide, x, 2.80, 2.70, 2.48, WHITE, line=LINE, radius=True)
        text(slide, number, x + 0.26, 3.08, 0.52, 0.32, size=20,
             shade=accent, bold=True, font=DISPLAY)
        text(slide, title, x + 0.26, 3.68, 2.08, 0.33, size=16,
             shade=NAVY, bold=True)
        text(slide, detail, x + 0.26, 4.18, 2.12, 0.72, size=11,
             shade=SLATE)
        if i < 3:
            line(slide, x + 2.70, 4.03, x + 3.10, 4.03, accent, 2)
    text(slide, "Store-scoped records keep daily activity aligned to the selected store.",
         0.78, 5.74, 11.6, 0.36, size=13, shade=SLATE, bold=True,
         align=PP_ALIGN.CENTER)
    footer(slide, 4)

    # 5. Sales and finance
    slide = base_slide(pres)
    heading(slide, "Sales and finance", "From daily entry to monthly performance.",
            "Sales, expenses, withdrawals, payroll, and profit are presented in related store workflows.")
    rect(slide, 0.72, 2.45, 4.03, 3.68, NAVY, radius=True)
    text(slide, "DAILY SALES", 1.04, 2.82, 3.1, 0.25, size=9,
         shade="B7C6D3", bold=True)
    text(slide, "One daily record", 1.04, 3.30, 3.1, 0.4, size=19,
         shade=WHITE, bold=True, font=DISPLAY)
    sales_rows = [("Sales total", TEAL), ("Payment channels", ORANGE),
                  ("Expenses", YELLOW), ("Cash handover", TEAL)]
    for i, (name, accent) in enumerate(sales_rows):
        yy = 4.02 + i * 0.43
        rect(slide, 1.06, yy + 0.08, 0.10, 0.10, accent, radius=True)
        text(slide, name, 1.30, yy, 2.72, 0.27, size=11, shade="E0E8ED")
    text(slide, "P&L", 5.34, 2.61, 1.3, 0.35, size=16, shade=ORANGE, bold=True)
    text(slide, "A month-by-month view", 5.34, 3.04, 3.7, 0.32,
         size=18, shade=NAVY, bold=True, font=DISPLAY)
    finance = [("Revenue", "Sales activity"), ("Store expenses", "Operating costs"),
               ("Payroll", "Salary projection"), ("Net profit", "Outcome and margin")]
    for i, (name, detail) in enumerate(finance):
        yy = 3.74 + i * 0.56
        rect(slide, 5.36, yy, 3.43, 0.43, WHITE, line=LINE, radius=True)
        text(slide, name, 5.58, yy + 0.10, 1.63, 0.19, size=10, shade=NAVY, bold=True)
        text(slide, detail, 7.22, yy + 0.10, 1.35, 0.19, size=8, shade=SLATE,
             align=PP_ALIGN.RIGHT)
    line(slide, 9.30, 2.75, 9.30, 5.83, LINE, 1)
    text(slide, "Related workflows", 9.72, 2.76, 2.6, 0.32, size=14,
         shade=NAVY, bold=True)
    add_line_items(slide, [
        ("Sales", "Daily record and review"),
        ("Withdrawals", "Track store cash movements"),
        ("P&L", "Revenue, costs, and margin"),
    ], 9.74, 3.36, 2.62, accent=ORANGE, gap=0.72)
    footer(slide, 5)

    # 6. People operations
    slide = base_slide(pres)
    heading(slide, "People operations", "A connected view of the employee lifecycle.",
            "Employee records support attendance, salary administration, and employee documentation.")
    stages = [
        ("EMPLOYEES", "Profiles, joining dates,\nemployment details", TEAL),
        ("ATTENDANCE", "Daily and weekly\nattendance workflows", ORANGE),
        ("SALARY", "Monthly projections,\nincentives, advances", YELLOW),
        ("SLIPS", "Downloadable salary\nslip PDFs", TEAL),
    ]
    for i, (title, detail, accent) in enumerate(stages):
        x = 0.76 + i * 3.13
        rect(slide, x, 2.82, 2.65, 2.06, WHITE, line=LINE, radius=True)
        rect(slide, x, 2.82, 2.65, 0.08, accent)
        text(slide, title, x + 0.24, 3.18, 2.18, 0.27, size=11,
             shade=accent, bold=True)
        text(slide, detail, x + 0.24, 3.69, 2.18, 0.72, size=13,
             shade=NAVY, bold=True)
        if i < 3:
            line(slide, x + 2.65, 3.84, x + 3.08, 3.84, accent, 2)
    rect(slide, 0.76, 5.36, 11.98, 0.70, MIST, radius=True)
    text(slide, "Store admins work with the employee records and salary activity for their authorized store.",
         1.03, 5.58, 11.4, 0.25, size=12, shade=NAVY, bold=True,
         align=PP_ALIGN.CENTER)
    footer(slide, 6)

    # 7. Inventory and configuration
    slide = base_slide(pres)
    heading(slide, "Inventory and store setup", "Standardize the catalog. Track stock by store.",
            "A shared product list sits alongside store-specific quantities and expiry dates.")
    rect(slide, 0.74, 2.48, 5.72, 3.75, WHITE, line=LINE, radius=True)
    label(slide, "Shared catalog", 1.06, 2.82, 1.62, fill=MIST, shade=TEAL)
    text(slide, "Product definitions stay consistent.", 1.06, 3.38, 4.92, 0.39,
         size=17, shade=NAVY, bold=True, font=DISPLAY)
    add_line_items(slide, [
        ("Products", "One catalog available to stores"),
        ("Stock batches", "Quantity and expiry by store"),
        ("Expiry control", "FEFO workflow and expired-stock disposal"),
    ], 1.10, 4.08, 4.90, accent=TEAL, gap=0.58)
    rect(slide, 6.82, 2.48, 5.72, 3.75, WHITE, line=LINE, radius=True)
    label(slide, "Store configuration", 7.14, 2.82, 2.02, fill=PEACH, shade=ORANGE)
    text(slide, "Local rules for daily execution.", 7.14, 3.38, 4.92, 0.39,
         size=17, shade=NAVY, bold=True, font=DISPLAY)
    add_line_items(slide, [
        ("Branding", "Store name and logo"),
        ("Salary cycle", "Calendar or custom monthly cycle"),
        ("Shifts and positions", "Store-defined roster and employee options"),
    ], 7.18, 4.08, 4.90, accent=ORANGE, gap=0.58)
    footer(slide, 7)

    # 8. Application screenshots
    slide = base_slide(pres)
    heading(slide, "Application screenshots", "Two views. One operating platform.",
            "Replace each frame with a 13-inch MacBook Pro desktop capture. Target: 1440 × 900 CSS pixels (2880 × 1800 Retina).")
    screenshot_frames = [
        (0.72, "01", "ADMIN DASHBOARD", "Network overview, stores, and users", TEAL),
        (6.84, "02", "STORE ADMIN DASHBOARD", "Store-level operations and performance", ORANGE),
    ]
    for x, number, title, detail, accent in screenshot_frames:
        rect(slide, x, 2.42, 5.76, 3.60, "E8EFED", line=accent, radius=True)
        rect(slide, x + 0.08, 2.50, 5.60, 3.44, "F5F7F4", line=LINE, radius=True)
        rect(slide, x + 2.63, 3.78, 0.42, 0.32, accent, radius=True)
        rect(slide, x + 2.72, 3.87, 0.24, 0.14, WHITE, radius=True)
        text(slide, "INSERT SCREENSHOT", x + 0.35, 4.36, 5.06, 0.27,
             size=12, shade=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text(slide, "16:10 desktop capture · fit to frame", x + 0.35, 4.75,
             5.06, 0.20, size=9, shade=SLATE, align=PP_ALIGN.CENTER)
        text(slide, number, x, 6.20, 0.42, 0.20, size=9,
             shade=accent, bold=True)
        text(slide, title, x + 0.47, 6.20, 3.20, 0.20, size=9,
             shade=NAVY, bold=True)
        text(slide, detail, x + 0.47, 6.48, 5.10, 0.18,
             size=9, shade=SLATE)
    footer(slide, 8)

    # 9. Access and mobile
    slide = base_slide(pres)
    heading(slide, "Access and usability", "Designed for clear access on desktop and mobile.",
            "Role-aware navigation and responsive layouts support work at the store and at the center.")
    rect(slide, 0.72, 2.42, 6.10, 3.90, WHITE, line=LINE, radius=True)
    text(slide, "ACCESS MODEL", 1.05, 2.78, 4.9, 0.22,
         size=9, shade=TEAL, bold=True)
    add_line_items(slide, [
        ("Role-based screens", "Navigation follows the signed-in role"),
        ("Store authorization", "Store workflows are scoped to approved access"),
        ("Verification", "Login includes a verification-code step"),
        ("Sensitive employee fields", "Protected employee information in admin flows"),
    ], 1.08, 3.28, 5.24, accent=TEAL, gap=0.67)
    rect(slide, 7.34, 2.42, 2.18, 3.90, NAVY, radius=True)
    rect(slide, 7.51, 2.68, 1.84, 3.35, PAPER, radius=True)
    rect(slide, 7.70, 2.93, 1.46, 0.34, TEAL, radius=True)
    text(slide, "DASHBOARD", 7.82, 3.04, 1.22, 0.10,
         size=6, shade=WHITE, bold=True, align=PP_ALIGN.CENTER)
    for i, (name, accent) in enumerate([("Sales summary", TEAL),
                                        ("Team", ORANGE),
                                        ("Attendance", YELLOW)]):
        yy = 3.53 + i * 0.69
        rect(slide, 7.70, yy, 1.46, 0.54, WHITE, line=LINE, radius=True)
        rect(slide, 7.80, yy + 0.13, 0.12, 0.12, accent, radius=True)
        text(slide, name, 7.99, yy + 0.14, 1.03, 0.15,
             size=6, shade=NAVY, bold=True)
    rect(slide, 7.98, 5.70, 0.92, 0.05, LINE, radius=True)
    text(slide, "RESPONSIVE VIEW", 7.21, 6.46, 2.45, 0.18,
         size=8, shade=SLATE, bold=True, align=PP_ALIGN.CENTER)
    line(slide, 10.05, 2.75, 10.05, 5.96, LINE, 1)
    text(slide, "Mobile-ready patterns", 10.42, 2.82, 2.22, 0.56,
         size=15, shade=NAVY, bold=True, font=DISPLAY)
    bullet_list(slide, [
        "Responsive page layouts",
        "Drawer-based entry forms",
        "Compact summary views",
    ], 10.42, 3.68, 2.26, line_height=0.62, size=10, accent=ORANGE)
    footer(slide, 9)

    # 10. Close
    slide = base_slide(pres, NAVY)
    label(slide, "Polar Bear  /  Application overview", 0.76, 0.72, 2.82,
          fill="243C54", shade="C8E9ED")
    text(slide, "A practical foundation\nfor consistent store work.",
         0.76, 1.50, 8.5, 1.34, size=30, shade=WHITE,
         bold=True, font=DISPLAY)
    text(slide, "One application connecting oversight, daily execution,\nand performance review.",
         0.80, 3.08, 7.8, 0.72, size=16, shade="D0DAE2")
    line(slide, 0.80, 4.27, 12.48, 4.27, "39536A", 0.8)
    text(slide, "DISCUSSION WITH OFFICIALS", 0.80, 4.72, 3.7, 0.24,
         size=9, shade=ORANGE, bold=True)
    prompts = [
        ("01", "Confirm store workflows and role ownership"),
        ("02", "Align reporting and payroll policies"),
        ("03", "Agree rollout and training priorities"),
    ]
    for i, (number, prompt) in enumerate(prompts):
        yy = 5.18 + i * 0.42
        text(slide, number, 0.82, yy, 0.40, 0.22, size=10,
             shade=TEAL, bold=True)
        text(slide, prompt, 1.35, yy, 7.7, 0.25,
             size=12, shade=WHITE)
    rect(slide, 10.42, 1.80, 1.62, 1.62, TEAL, radius=True)
    text(slide, "PB", 10.42, 2.32, 1.62, 0.42, size=30,
         shade=WHITE, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER)
    text(slide, "STORE OPERATIONS", 9.45, 3.73, 3.55, 0.24,
         size=9, shade="B7C6D3", bold=True, align=PP_ALIGN.CENTER)
    footer(slide, 10, dark=True)

    pres.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = make_deck()
    print(f"Created {path}")
